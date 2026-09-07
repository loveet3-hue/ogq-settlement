# -*- coding: utf-8 -*-
"""네이버 출신 크리에이터 타 마켓 판매 로우 데이터 -> 마켓별 분류 + 네이버 수수료 추합본 생성.

수수료 구조
    정산 대상 금액 = 판매액 - 결제 수수료 - 마켓 수수료
    네이버 수수료  = 정산 대상 금액 * 15%
따라서 판매액 대비 실효율 = (1 - 결제수수료율 - 마켓수수료율) * 0.15
"""
import csv
import io
import re
import collections
from copy import copy
from datetime import datetime

import openpyxl
from decimal import Decimal, ROUND_HALF_EVEN

import tables

NAVER_SHARE = 0.15
INFO_SHEET = "수수료 안내"
# 0이면 '-'로 보이게 (회사 계산기와 같은 표기). 값은 0 그대로라 수식에는 영향 없음.
FMT_MONEY = '#,##0.00;-#,##0.00;"-"'
FMT_RATE = '0.000%;-0.000%;"-"'
FMT_WON = '#,##0;-#,##0;"-"' 


def won(x):
    """원 단위 반올림. 정확히 0.5인 경우는 짝수 쪽으로 붙인다.

    6,000 × 0.14175 = 850.5 → 850,  10,000 × 0.14175 = 1,417.5 → 1,418
    """
    d = x if isinstance(x, Decimal) else Decimal(repr(float(x)))
    return int(d.quantize(Decimal("1"), rounding=ROUND_HALF_EVEN))


def round_even_formula(expr):
    """엑셀에는 짝수쪽 반올림이 없어 직접 만든다.

    안쪽 ROUND(...,6)은 부동소수점 찌꺼기 제거용
    (없으면 6,000×0.14175 가 850.4999999999999 로 계산된다).
    정확히 .5면 2*ROUND(x/2,0) 이 짝수 쪽으로 붙고, 아니면 보통 반올림.
    """
    e = "ROUND(%s,6)" % expr
    return "IF(MOD({e}*2,2)=1,2*ROUND({e}/2,0),ROUND({e},0))".format(e=e)


def settle(total, market):
    """판매액 -> 단계별 수수료. 회사 정산 계산기와 같은 방식이다.

        결제 수수료   = 판매액 × 결제율
        마켓 수수료   = 판매액 × 마켓율
        정산 대상 금액 = 판매액 − 결제 수수료 − 마켓 수수료
        네이버 수수료  = ROUND(정산 대상 금액 × 15%, 0)

    중간 단계는 반올림하지 않는다. 마지막 네이버 수수료만 원 단위로 반올림한다.
    (10,000원 SOOP: 마켓 1,417.5 / 정산 대상 8,032.5 / 네이버 1,205)
    전 과정을 Decimal로 계산해 6,000 × 0.14175 가 850.4999…로 새는 것을 막는다.
    """
    d = Decimal(total)
    pay = d * Decimal(repr(market.pay_fee))
    mk = d * Decimal(repr(market.market_fee))
    base = d - pay - mk
    return dict(total=total, pay=pay, market_fee=mk, base=base,
                naver=won(base * Decimal(repr(NAVER_SHARE))))
HEADERS = ["거래시간", "거래 ID", "거래유형", "크리에이터ID", "판매마켓 ID",
           "상품 코드", "상품명", "가격"]
COLS = "ABCDEFGHI"
FULL_HEADERS = ["거래시간", "거래 ID", "거래유형", "크리에이터ID", "판매마켓 ID",
                "판매마켓", "상품 코드", "상품명", "가격"]


class Market:
    """마켓 1개의 수수료 구조와 추합본 시트 규칙."""

    def __init__(self, key, label, pay_fee, market_fee, book, cycle,
                 offset=1, month_fmt="{m}월분"):
        self.key = key                # CSV의 '판매마켓' 값
        self.label = label            # 화면 표시명
        self.pay_fee = pay_fee        # 결제 수수료율
        self.market_fee = market_fee  # 마켓 수수료율
        self.book = book              # 추합본 파일 구분명
        self.cycle = cycle            # "month" | "quarter"
        self.offset = offset          # 정산월 = 판매월 + offset
        self.month_fmt = month_fmt

    @property
    def settle_ratio(self):
        return 1 - self.pay_fee - self.market_fee

    @property
    def rate(self):
        """판매액 대비 네이버 수수료 실효율."""
        return round(self.settle_ratio * NAVER_SHARE, 10)

    @property
    def filename(self):
        return "추합본_%s_네이버수수료.xlsx" % self.book

    def sheet_name(self, dates):
        months = sorted({d.month for d in dates})
        year = min(dates).year
        if self.cycle == "quarter":
            q = (months[0] - 1) // 3 + 1
            return "%d분기 정산(%d.%d월~%d월 분)" % (q, year, (q - 1) * 3 + 1, q * 3)
        m = months[0]
        settle, sy = m + self.offset, year
        if settle > 12:
            settle, sy = settle - 12, year + 1
        return "%d월정산(%d.%s)" % (settle, sy, self.month_fmt.format(m=m))


MARKETS = [
    Market("SOOP OGQ 마켓",  "SOOP OGQ이모티콘", 0.055, 0.14175, "SOOP",       "month",   offset=1),
    Market("채팅+ 원스토어",  "원스토어 (RCS)", 0.0,   0.40,    "RCS",        "month",   offset=2),
    Market("연합뉴스",        "연합뉴스",     0.0,   0.40,    "연합",       "month",   offset=1),
    Market("채팅+ OGQ 마켓",  "채팅+ OGQ마켓", 0.055, 0.20,    "채팅플러스", "quarter"),
    Market("OGQ 그라폴리오",  "그라폴리오",   0.055, 0.0,     "그라폴리오", "month",   offset=1, month_fmt="{m:02d}월분"),
    Market("네이버 밴드",     "밴드",         0.0,   0.30,    "밴드",       "month",   offset=2),
    Market("프리미엄뉴스",    "프리미엄뉴스", 0.0,   0.0,     "프리미엄뉴스", "month", offset=2),
]
BY_KEY = {m.key: m for m in MARKETS}

DATE_FORMATS = ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M",
                "%Y.%m.%d %H:%M:%S", "%Y.%m.%d %H:%M",
                "%Y/%m/%d %H:%M:%S", "%Y/%m/%d %H:%M")


def parse_dt(value):
    if isinstance(value, datetime):
        return value
    s = str(value).strip()
    for f in DATE_FORMATS:
        try:
            return datetime.strptime(s, f)
        except ValueError:
            pass
    return s


def read_csv(data, filename=""):
    """CSV · 엑셀 · 넘버스 로우 데이터를 읽어 행 리스트로 돌려준다."""
    sheets = tables.load_sheets(data, filename)
    found = tables.find_table(sheets, FULL_HEADERS)
    if not found:
        have = ", ".join(tables.text(c) for c in (sheets[0][1][0] if sheets and sheets[0][1] else []))
        raise ValueError("필수 컬럼(%s)을 가진 표를 찾지 못했습니다. 첫 줄: %s"
                         % (" · ".join(FULL_HEADERS), have or "(빈 파일)"))
    _name, _header, rows, idx = found
    out = []
    for r in rows:
        if len(r) <= max(idx.values()):
            continue
        cell = {k: r[i] for k, i in idx.items()}
        if cell["거래 ID"] in (None, "") or cell["가격"] in (None, ""):
            continue
        out.append([
            parse_dt(cell["거래시간"]),
            int(float(tables.text(cell["거래 ID"]))),
            tables.text(cell["거래유형"]),
            tables.text(cell["크리에이터ID"]),
            tables.text(cell["판매마켓 ID"]),
            tables.text(cell["판매마켓"]),
            tables.text(cell["상품 코드"]),
            tables.text(cell["상품명"]),
            int(float(tables.text(cell["가격"]))),
        ])
    if not out:
        raise ValueError("읽을 수 있는 거래가 없습니다.")
    return out


def classify(records):
    """판매마켓별로 묶어 {마켓명: [row,...]} 반환. 미등록 마켓도 그대로 포함."""
    buckets = collections.defaultdict(list)
    for r in records:
        buckets[r[5]].append(r)
    for rows in buckets.values():
        rows.sort(key=lambda r: (r[0] if isinstance(r[0], datetime) else datetime.max, r[1]))
    return dict(buckets)


# ---------------------------------------------------------------- 시트 서식

_ORANGE = "FFFFC000"
_BLUE = "FF305496"


def _style_from(ws_src, ws_dst):
    """기존 추합본 시트의 서식을 그대로 복제."""
    for col, dim in ws_src.column_dimensions.items():
        ws_dst.column_dimensions[col].width = dim.width
    for col in COLS:
        for row in (1, 2, 3):
            src = ws_src["%s%d" % (col, row)]
            ws_dst["%s%d" % (col, row)]._style = copy(src._style)


def _style_default(ws):
    """참조할 기존 파일이 없을 때 쓰는 기본 서식."""
    from openpyxl.styles import Font, PatternFill, Alignment
    widths = dict(A=18.56, B=9.56, C=11.44, D=14.33, E=13.56, F=13.33, G=12.67, H=23.22, I=6.44)
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    for col in COLS:
        c1, c2 = ws["%s1" % col], ws["%s2" % col]
        c1.fill = PatternFill("solid", fgColor=_ORANGE)
        c1.font = Font(bold=True)
        c1.alignment = Alignment(horizontal="center")
        c2.fill = PatternFill("solid", fgColor=_BLUE)
        c2.font = Font(bold=True, color="FFFFFFFF")
        c2.alignment = Alignment(horizontal="center")
    ws["B1"].number_format = "#,##0"
    ws["A3"].number_format = "yyyy\\-mm\\-dd\\ hh:mm:ss"
    ws["B3"].number_format = "0"
    ws["I3"].number_format = "#,##0;\\-#,##0"


def _read_sheet_rows(ws):
    rows = []
    for r in ws.iter_rows(min_row=3, values_only=True):
        if len(r) < 9 or r[1] is None:
            continue
        rows.append([parse_dt(r[0]), int(r[1]), r[2], r[3], r[4], r[5], r[6], r[7], int(r[8] or 0)])
    return rows


def write_sheet(wb, market, rows, tmpl=None, merge=True):
    """wb 맨 앞에 정산 시트를 만든다. 같은 이름 시트가 있으면 거래 ID 기준 병합."""
    name = market.sheet_name([r[0] for r in rows if isinstance(r[0], datetime)])
    merged_from = 0
    if name in wb.sheetnames:
        old = wb[name]
        if merge:
            existing = _read_sheet_rows(old)
            seen = {r[1] for r in rows}
            keep = [r for r in existing if r[1] not in seen]
            merged_from = len(keep)
            rows = keep + rows
            rows.sort(key=lambda r: (r[0] if isinstance(r[0], datetime) else datetime.max, r[1]))
        if tmpl is None:
            tmpl = old
        del wb[name]

    ws = wb.create_sheet(name, 0)
    last = len(rows) + 2
    total_ref = "SUM(I3:I%d)" % last
    ws["A1"] = "네이버 수수료"
    ws["B1"] = "=" + round_even_formula("%s*%s" % (total_ref, repr(market.rate)))
    ws["D1"] = ("판매액에서 결제 수수료 %s, 마켓 수수료 %s를 뺀 정산 대상 금액의 15%%  "
                "(판매액 대비 %s)  ·  계산 과정은 '%s' 시트"
                % (_pct(market.pay_fee), _pct(market.market_fee, 3),
                   _pct(market.rate, 5), INFO_SHEET))
    for col, h in zip(COLS, FULL_HEADERS):
        ws["%s2" % col] = h
    for i, r in enumerate(rows, start=3):
        for j, col in enumerate(COLS):
            ws["%s%d" % (col, i)] = r[j]

    if tmpl is not None:
        _style_from(tmpl, ws)
    else:
        _style_default(ws)
    # 3행 서식을 데이터 전체에 확장
    for i in range(4, last + 1):
        for col in COLS:
            ws["%s%d" % (col, i)]._style = copy(ws["%s3" % col]._style)

    total = sum(r[8] for r in rows)
    calc = settle(total, market)
    return dict(sheet=name, rows=len(rows), merged=merged_from,
                total=total, fee=calc["naver"], calc=calc)


def _pct(v, digits=2):
    return "-" if not v else ("%.*f%%" % (digits, v * 100)).rstrip()


def write_info_sheet(wb, market=None, sheet_name=None, last_row=None, note=None):
    """파일 설명 · 수수료 계산 과정 · 마켓별 요율을 담은 안내 시트를 맨 뒤에 만든다.

    market 과 sheet_name 을 주면 그 정산 시트를 참조하는 살아있는 계산식을 넣는다.
    market 이 None 이면 전체 합본용 안내가 된다.
    """
    from openpyxl.styles import Alignment, Font, PatternFill

    if INFO_SHEET in wb.sheetnames:
        del wb[INFO_SHEET]
    ws = wb.create_sheet(INFO_SHEET)
    for col, w in zip("ABCDEF", (26, 14, 18, 18, 16, 20)):
        ws.column_dimensions[col].width = w

    head = Font(bold=True, size=13)
    bold = Font(bold=True)
    label = Font(bold=True, color="FF44546A", size=11)
    band = PatternFill("solid", fgColor="FF44546A")
    mine = PatternFill("solid", fgColor="FFFFF2CC")
    money = FMT_MONEY

    def put(r, c, v, font=None, fill=None, fmt=None, align=None):
        cell = ws.cell(row=r, column=c, value=v)
        if font:
            cell.font = font
        if fill:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        if align:
            cell.alignment = Alignment(horizontal=align)
        return cell

    put(1, 1, "OGQ 정산 — 네이버 출신 작가 정산", head)
    put(2, 1, "만든 곳: https://ogq-settlement.streamlit.app")

    put(4, 1, "이 파일", label)
    info = ([("판매마켓", market.label),
             ("로우 데이터 표기", market.key),
             ("추합본 파일", market.filename),
             ("시트 주기", "분기별 (한 시트에 3개월)" if market.cycle == "quarter"
                           else "월별 — 정산월 = 판매월 + %d" % market.offset)]
            if market else
            [("파일 종류", "전체 합본 — 마켓별 전체 내역 + 정산 요약")]
            + ([("포함 마켓", note)] if note else []))
    info.append(("만든 날짜", datetime.now().strftime("%Y-%m-%d")))
    for i, (k, v) in enumerate(info):
        put(5 + i, 1, k, bold)
        put(5 + i, 2, v)

    r = 5 + len(info) + 1
    put(r, 1, "수수료를 어떻게 떼는가", label)
    put(r + 1, 1, "판매액에서 ① 결제 수수료와 ② 마켓 수수료를 뺀 금액이 '정산 대상 금액'이고,")
    put(r + 2, 1, "그 금액의 15%를 네이버가 가져갑니다. 이 15%가 정산금입니다.")
    put(r + 3, 1, "중간 단계는 반올림하지 않고, 마지막 정산금만 원 단위로 반올림합니다.")

    r += 5
    if market and sheet_name and last_row:
        put(r, 1, "이 파일의 계산  (시트: %s)" % sheet_name, label)
        h = r + 1
        for i, name in enumerate(["항목", "수수료율", "금액"]):
            put(h, 1 + i, name, Font(bold=True, color="FFFFFFFF"), band, None, "center")
        ref = "'%s'!I3:I%d" % (sheet_name, last_row)
        put(h + 1, 1, "판매액 합계", bold)
        put(h + 1, 3, "=SUM(%s)" % ref, None, None, money)
        put(h + 2, 1, "① 결제 수수료")
        put(h + 2, 2, market.pay_fee, None, None, FMT_RATE)
        put(h + 2, 3, "=C%d*B%d" % (h + 1, h + 2), None, None, money)
        put(h + 3, 1, "② 마켓 수수료")
        put(h + 3, 2, market.market_fee, None, None, FMT_RATE)
        put(h + 3, 3, "=C%d*B%d" % (h + 1, h + 3), None, None, money)
        put(h + 4, 1, "정산 대상 금액 = 판매액 − ① − ②", bold)
        put(h + 4, 3, "=C%d-C%d-C%d" % (h + 1, h + 2, h + 3), bold, None, money)
        put(h + 5, 1, "네이버 수수료 (정산금)", bold, mine)
        put(h + 5, 2, NAVER_SHARE, None, mine, "0%")
        put(h + 5, 3, "=" + round_even_formula("C%d*B%d" % (h + 4, h + 5)),
            bold, mine, "#,##0")
        put(h + 6, 1, "판매액 대비 실효율")
        put(h + 6, 2, "=IF(C%d=0,0,C%d/C%d)" % (h + 1, h + 5, h + 1), None, None, "0.00000%")
        put(h + 7, 1, "→ 정산 시트 B1 셀의 값과 같습니다.")
        r = h + 9

    put(r, 1, "마켓별 수수료율", label)
    for i, name in enumerate(["판매마켓", "① 결제 수수료", "② 마켓 수수료",
                              "정산 대상 비율", "네이버 몫", "판매액 대비 실효율"]):
        put(r + 1, 1 + i, name, Font(bold=True, color="FFFFFFFF"), band, None, "center")
    for i, m in enumerate(MARKETS):
        rr = r + 2 + i
        fill = mine if (market and m.key == market.key) else None
        put(rr, 1, m.label, bold if fill else None, fill)
        put(rr, 2, m.pay_fee, None, fill, FMT_RATE)
        put(rr, 3, m.market_fee, None, fill, FMT_RATE)
        put(rr, 4, m.settle_ratio, None, fill, "0.000%")
        put(rr, 5, NAVER_SHARE, None, fill, "0%")
        put(rr, 6, m.rate, bold if fill else None, fill, "0.00000%")
    r += 2 + len(MARKETS)
    if market:
        put(r, 1, "※ 노란색이 이 파일의 마켓입니다.")
        r += 1

    put(r + 1, 1, "참고", label)
    for i, line in enumerate([
            "REFUNDMENT(환불)는 가격이 음수로 들어와 판매액 합계에서 자동으로 차감됩니다.",
            "새 정산분은 항상 맨 앞 시트로 추가되고, 과거 시트와 서식은 그대로 둡니다.",
            "채팅+ OGQ마켓만 분기 단위 시트라 같은 분기 시트에 이어붙입니다 (거래 ID로 중복 제거).",
            "수수료율이 바뀌면 이 파일이 아니라 만든 곳에서 고쳐야 다음 달부터 반영됩니다.",
    ]):
        put(r + 2 + i, 1, "· " + line)
    return ws


def detect_market(wb):
    """추합본 워크북의 첫 시트를 보고 어느 마켓 파일인지 판별."""
    counts = collections.Counter()
    for ws in wb.worksheets:
        for r in ws.iter_rows(min_row=3, max_row=40, values_only=True):
            if len(r) > 5 and r[5]:
                counts[r[5]] += 1
        if counts:
            break
    if counts:
        key = counts.most_common(1)[0][0]
        if key in BY_KEY:
            return BY_KEY[key]
    return None


def build(records, existing=None, merge=True):
    """records + {마켓명: 기존 워크북 bytes} -> [{market, filename, bytes, ...}]

    existing 에 없는 마켓은 동일 양식으로 새 파일을 만든다.
    """
    existing = existing or {}
    buckets = classify(records)
    results, unknown = [], []

    for name, rows in sorted(buckets.items(), key=lambda kv: -sum(r[8] for r in kv[1])):
        market = BY_KEY.get(name)
        if market is None:
            unknown.append(dict(market=name, rows=len(rows),
                                total=sum(r[8] for r in rows)))
            continue
        blob = existing.get(name)
        if blob:
            wb = openpyxl.load_workbook(io.BytesIO(blob))
            tmpl = wb.worksheets[0]
        else:
            wb = openpyxl.Workbook()
            wb.remove(wb.active)
            tmpl = None
        info = write_sheet(wb, market, rows, tmpl=tmpl, merge=merge)
        write_info_sheet(wb, market, sheet_name=info["sheet"],
                         last_row=len(rows) + 2)
        buf = io.BytesIO()
        wb.save(buf)
        info.update(market=market.label, key=market.key, rate=market.rate,
                    filename=market.filename, bytes=buf.getvalue(),
                    is_new=not blob, sheets=wb.sheetnames,
                    records=rows, obj=market)
        results.append(info)

    return results, unknown


def breakdown(amount, market):
    """계산기용: 판매액 -> 각 수수료 금액 (settle과 같은 순서)."""
    c = settle(amount, market)
    return dict(pay=c["pay"], market=c["market_fee"], base=c["base"], naver=c["naver"])


def build_combined(results, title="네이버 출신 작가 정산"):
    """마켓별 결과 -> 전체 내역 + 정산 요약을 담은 합본 워크북 bytes.

    요약의 금액은 모두 수식이라 내역 시트를 고치면 따라 바뀐다.
    """
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("정산 요약")

    band = PatternFill("solid", fgColor="FF44546A")
    yellow = PatternFill("solid", fgColor="FFFFFF00")
    white = Font(bold=True, color="FFFFFFFF")
    bold = Font(bold=True)
    red = Font(bold=True, color="FFC00000")
    money = FMT_MONEY

    for col, w in zip("BCDEFGHIJK",
                      (20, 24, 16, 13, 16, 13, 16, 17, 11, 15)):
        ws.column_dimensions[col].width = w

    ws["B1"] = title + " — 전체 합본"
    ws["B1"].font = Font(bold=True, size=13)
    ws["B2"] = ("판매액 − 결제 수수료 − 마켓 수수료 = 정산 대상 금액,"
                " 그 금액의 15%가 정산금입니다.")

    headers = ["판매마켓", "정산 시트", "판매액", "① 결제 수수료율", "① 결제 수수료",
               "② 마켓 수수료율", "② 마켓 수수료", "정산 대상 금액",
               "네이버 몫", "정산금"]
    for i, name in enumerate(headers):
        c = ws.cell(row=4, column=2 + i, value=name)
        c.fill = band
        c.font = white
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    r0 = 5
    for i, res in enumerate(results):
        r = r0 + i
        m = res["obj"]
        sheet = m.label
        n = len(res["records"]) + 1
        ws.cell(row=r, column=2, value=m.label).font = bold
        ws.cell(row=r, column=3, value=res["sheet"])
        ws.cell(row=r, column=4, value="=SUM('%s'!I2:I%d)" % (sheet, n)).number_format = money
        ws.cell(row=r, column=5, value=m.pay_fee).number_format = FMT_RATE
        ws.cell(row=r, column=6, value="=D%d*E%d" % (r, r)).number_format = money
        ws.cell(row=r, column=7, value=m.market_fee).number_format = FMT_RATE
        ws.cell(row=r, column=8, value="=D%d*G%d" % (r, r)).number_format = money
        ws.cell(row=r, column=9,
                value="=D%d-F%d-H%d" % (r, r, r)).number_format = money
        ws.cell(row=r, column=10, value=NAVER_SHARE).number_format = "0%"
        c = ws.cell(row=r, column=11,
                    value="=" + round_even_formula("I%d*J%d" % (r, r)))
        c.number_format = FMT_WON
        c.font = red

    rt = r0 + len(results)
    ws.cell(row=rt, column=2, value="합계").font = bold
    for col in "DFHIK":
        c = ws.cell(row=rt, column=ord(col) - 64,
                    value="=SUM(%s%d:%s%d)" % (col, r0, col, rt - 1))
        c.number_format = FMT_WON if col == "K" else money
        c.font = red if col == "K" else bold
    for col in range(2, 12):
        ws.cell(row=rt, column=col).fill = yellow

    ws.cell(row=rt + 2, column=2,
            value="※ 판매액은 아래 마켓별 시트에서 자동으로 합산됩니다. "
                  "정산금만 원 단위로 반올림합니다.")

    # 마켓별 전체 내역
    for res in results:
        sh = wb.create_sheet(res["obj"].label)
        for col, w in zip(COLS, (18.6, 9.6, 11.4, 14.3, 13.6, 13.3, 12.7, 23.2, 9)):
            sh.column_dimensions[col].width = w
        for i, name in enumerate(FULL_HEADERS):
            c = sh.cell(row=1, column=1 + i, value=name)
            c.fill = PatternFill("solid", fgColor="FF305496")
            c.font = white
            c.alignment = Alignment(horizontal="center")
        for i, row in enumerate(res["records"], start=2):
            for j, v in enumerate(row):
                c = sh.cell(row=i, column=1 + j, value=v)
                if j == 0:
                    c.number_format = "yyyy\\-mm\\-dd\\ hh:mm:ss"
                elif j == 1:
                    c.number_format = "0"
                elif j == 8:
                    c.number_format = "#,##0;-#,##0"
        sh.freeze_panes = "A2"

    write_info_sheet(wb, None, note=", ".join(r["obj"].label for r in results))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
