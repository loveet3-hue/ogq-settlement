# -*- coding: utf-8 -*-
"""CSV · 엑셀(xlsx) · 넘버스(numbers)를 같은 모양으로 읽어 온다.

    load_sheets(bytes | 경로, 파일명) -> [(시트명, [[셀, ...], ...]), ...]

확장자를 못 믿는 경우(업로드 위젯이 이름을 안 주는 등)에도 내용으로 판별한다.
"""
import csv
import io
import os
import zipfile

CSV, XLSX, NUMBERS = ".csv", ".xlsx", ".numbers"


def text(v):
    """셀 값을 문자열로. 숫자로 읽힌 ID의 꼬리 '.0'을 없앤다."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _sniff(data):
    """내용으로 형식 판별."""
    if data[:2] != b"PK":
        return CSV
    try:
        names = zipfile.ZipFile(io.BytesIO(data)).namelist()
    except zipfile.BadZipFile:
        return CSV
    if any(n.startswith("xl/") for n in names):
        return XLSX
    if any(n.startswith("Index") or n.endswith(".iwa") for n in names):
        return NUMBERS
    return XLSX


def _decode(data):
    for enc in ("utf-8-sig", "cp949", "euc-kr", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def load_sheets(data, filename=""):
    if isinstance(data, str):
        filename = filename or data
        data = open(data, "rb").read()

    kind = _sniff(data)
    ext = os.path.splitext(filename or "")[1].lower()
    if kind != CSV and ext in (XLSX, NUMBERS):
        kind = ext                      # zip 두 종류는 확장자를 우선 신뢰

    if kind == CSV:
        rows = [list(r) for r in csv.reader(io.StringIO(_decode(data)))]
        return [(os.path.splitext(os.path.basename(filename))[0] or "sheet1", rows)]

    if kind == NUMBERS:
        try:
            from numbers_parser import Document
        except ImportError:             # pragma: no cover
            raise ValueError(
                "넘버스 파일을 읽으려면 numbers-parser가 필요합니다 "
                "(pip install numbers-parser). CSV나 엑셀로 저장해 올려도 됩니다.")
        import tempfile
        with tempfile.NamedTemporaryFile(suffix=NUMBERS, delete=False) as fp:
            fp.write(data)
            path = fp.name
        try:
            doc = Document(path)
            out = []
            for sheet in doc.sheets:
                for table in sheet.tables:
                    name = table.name or sheet.name
                    out.append((name, [list(r) for r in table.rows(values_only=True)]))
            return out
        finally:
            os.unlink(path)

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    return [(ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
            for ws in wb.worksheets]


def find_table(sheets, required, skip=()):
    """required 컬럼을 모두 가진 첫 표를 찾아 (시트명, 헤더, 데이터행, {컬럼: 인덱스}) 반환."""
    for name, rows in sheets:
        if name.strip() in skip or not rows:
            continue
        header = [text(c) for c in rows[0]]
        if all(c in header for c in required):
            return name, rows[0], rows[1:], {c: header.index(c) for c in required}
    return None
